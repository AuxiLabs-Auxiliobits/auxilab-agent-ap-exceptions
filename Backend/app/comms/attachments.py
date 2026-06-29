"""Email attachment generation for consolidated communications.

When many invoices route to one recipient, inlining a full section per invoice
makes an unreadable wall-of-text email. Instead we attach the invoice list as a
spreadsheet (the recipient can sort/filter/total it) and keep the email body
short. ``openpyxl`` ships with the backend, so this needs no extra dependency.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

# MIME type for a .xlsx workbook.
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass(frozen=True)
class EmailAttachment:
    """A file to attach to an outbound email. ``maintype``/``subtype`` are the
    MIME parts passed straight to ``EmailMessage.add_attachment``."""

    filename: str
    data: bytes
    maintype: str = "application"
    subtype: str = "octet-stream"


def safe_filename(name: str, *, fallback: str = "AP") -> str:
    """Slug a vendor/title into a filesystem- and header-safe filename stem."""
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", (name or "").strip()).strip("-")
    return cleaned[:60] or fallback


# A column spec: (header, item-key, numeric?, width). numeric columns are
# written as real numbers (currency-formatted) so the recipient can total them.
_Column = tuple[str, str, bool, int]

_INVOICE_COLUMNS: list[_Column] = [
    ("Invoice ID", "invoice_id", False, 16),
    ("Vendor", "vendor_name", False, 22),
    ("Amount", "amount", True, 14),
    ("Issue", "issue", False, 20),
    ("Days outstanding", "days_outstanding", False, 16),
    ("PO number", "po_number", False, 16),
    ("Subject", "subject", False, 32),
    ("Message", "body", False, 60),
]

_SLA_COLUMNS: list[_Column] = [
    ("Invoice ID", "invoice_id", False, 16),
    ("Vendor", "vendor_name", False, 22),
    ("Amount", "invoice_amount", True, 14),
    ("Status", "status", False, 12),
    # "Recommended approach" = the resolution path as the operator-facing label
    # (derived in build_sla_digest_attachment), matching the digest email body.
    ("Recommended approach", "recommended_approach", False, 24),
    ("SLA hours", "sla_hours", False, 11),
    ("Hours overdue", "hours_overdue", False, 13),
    ("Comms sent", "comms_sent", False, 11),
    ("Priority", "priority", False, 10),
]


def _table_to_xlsx(*, sheet_title: str, columns: list[_Column], items: list[dict[str, Any]]) -> bytes:
    """Generic styled-table → .xlsx (bytes): bold header, frozen header row,
    autofilter, sized columns, numeric columns as real currency-formatted numbers."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = (sheet_title or "Sheet1")[:31]  # Excel caps sheet names at 31 chars

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="374151")
    ws.append([h for h, _, _, _ in columns])
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for it in items:
        row: list[Any] = []
        for _, key, numeric, _w in columns:
            val = it.get(key)
            if numeric:
                row.append(float(val) if val is not None else None)
            elif isinstance(val, bool):
                row.append("Yes" if val else "No")
            else:
                row.append("" if val is None else str(val))
        ws.append(row)

    for i, (_, _, numeric, w) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        if numeric:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=i).number_format = '#,##0.00'

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{ws.max_row}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_invoice_xlsx(items: list[dict[str, Any]], *, sheet_title: str = "Invoices") -> bytes:
    """Render the consolidated invoice list as an .xlsx workbook (bytes).

    ``items`` are the dicts assembled in ``dispatch_consolidated`` (keys:
    invoice_id, vendor_name, amount, issue, days_outstanding, po_number,
    subject, body)."""
    return _table_to_xlsx(sheet_title=sheet_title, columns=_INVOICE_COLUMNS, items=items)


def build_sla_digest_attachment(report: dict, *, run_id: str) -> EmailAttachment:
    """Build an .xlsx of the FULL SLA list (every breached + due-soon item) for
    the operator digest — so the email body stays a short summary and the
    complete list rides along as a spreadsheet."""
    from app.schemas import resolution_path_label

    raw = list(report.get("breached_items") or []) + list(report.get("due_soon_items") or [])
    # Surface the resolution path as the friendly "Recommended approach" label
    # (matches the digest email body + reviewer console), not the raw routing code.
    items = [
        {**it, "recommended_approach": resolution_path_label(it.get("resolution_path"))}
        for it in raw
    ]
    data = _table_to_xlsx(sheet_title="SLA digest", columns=_SLA_COLUMNS, items=items)
    return EmailAttachment(
        filename=f"sla-digest-{safe_filename(run_id[:8])}-{len(items)}.xlsx",
        data=data, maintype="application", subtype=XLSX_MIMETYPE.split("/", 1)[1],
    )


def build_invoice_attachment(
    items: list[dict[str, Any]], *, vendor: str | None, count: int
) -> EmailAttachment:
    """Convenience: build the .xlsx and wrap it as an EmailAttachment with a
    descriptive filename like ``invoices-Acme-78.xlsx``."""
    stem = safe_filename(vendor or "AP")
    filename = f"invoices-{stem}-{count}.xlsx"
    data = build_invoice_xlsx(items, sheet_title=f"{stem} invoices"[:31])
    return EmailAttachment(filename=filename, data=data, maintype="application",
                           subtype=XLSX_MIMETYPE.split("/", 1)[1])

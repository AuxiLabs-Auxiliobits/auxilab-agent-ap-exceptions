"""Large consolidated emails attach an .xlsx instead of an inline wall of text."""
from __future__ import annotations

import io
from decimal import Decimal


def _items(n):
    return [
        {
            "invoice_id": f"INV-{i}", "vendor_name": "Acme", "amount": float(1000 + i),
            "issue": "Price Variance", "days_outstanding": i, "po_number": f"PO-{i}",
            "subject": f"Invoice INV-{i}", "body": f"Please review INV-{i}.",
        }
        for i in range(1, n + 1)
    ]


def test_build_invoice_xlsx_has_header_and_rows():
    from openpyxl import load_workbook
    from app.comms.attachments import build_invoice_xlsx

    data = build_invoice_xlsx(_items(3))
    ws = load_workbook(io.BytesIO(data)).active
    assert ws.cell(row=1, column=1).value == "Invoice ID"
    assert ws.max_row == 4  # header + 3 rows
    assert ws.cell(row=2, column=1).value == "INV-1"
    # Amount is a real number (so the recipient can total it), not text.
    assert ws.cell(row=2, column=3).value == 1001.0


def test_render_consolidated_compact_omits_per_invoice_bodies():
    from app.comms.email_template import render_consolidated
    from app.config import Settings
    from app.schemas import CommChannel

    items = _items(12)
    note = "Full details for all 12 invoices are attached: invoices-Acme-12.xlsx"
    text, html = render_consolidated(
        channel=CommChannel.VENDOR_EMAIL, items=items, settings=Settings(), attachment_note=note,
    )
    # Compact: the note + count are present, the per-invoice wall of text is not.
    assert note in text and "12 invoice" in text.lower()
    assert "Please review INV-1." not in text
    assert "Please review INV-1." not in html
    assert "attached" in html.lower()


def _dispatch(monkeypatch, tmp_path, n):
    monkeypatch.setenv("COMMS_DRYRUN", "true")
    monkeypatch.setenv("COMMS_SENT_DIR", str(tmp_path / "sent"))
    monkeypatch.setenv("AUTH_ENABLED", "false")
    from app.config import get_settings
    get_settings.cache_clear()

    from app.comms.dispatcher import dispatch_consolidated
    from app.schemas import CommChannel, CommunicationDraft, ExceptionRow, SendStatus

    drafts, rows = [], {}
    for i in range(1, n + 1):
        inv = f"INV-{i}"
        drafts.append(CommunicationDraft(
            invoice_id=inv, channel=CommChannel.VENDOR_EMAIL, recipient_hint="Vendor AR",
            subject=f"Invoice {inv}", body=f"Please review {inv}.",
            template_id="vendor_price_variance", model_id="m",
        ))
        rows[inv] = ExceptionRow(
            invoice_id=inv, vendor_name="Acme", invoice_amount=Decimal("1000"),
            po_number="PO", exception_type="Price Variance", exception_description="x",
            days_outstanding=5,
        )
    res = dispatch_consolidated(
        recipient="ar@vendor.com", channel=CommChannel.VENDOR_EMAIL,
        drafts=drafts, rows_by_id=rows, run_id="run_att",
    )
    assert res.status == SendStatus.DRYRUN
    get_settings.cache_clear()
    return tmp_path / "sent" / "run_att"


def test_over_threshold_writes_xlsx_attachment(monkeypatch, tmp_path):
    out = _dispatch(monkeypatch, tmp_path, n=12)  # > default threshold of 10
    xlsx = list(out.glob("*.xlsx"))
    assert len(xlsx) == 1
    assert xlsx[0].name == "invoices-Acme-12.xlsx"
    assert list(out.glob("consolidated__*.eml"))  # the email itself still written


def test_under_threshold_stays_inline(monkeypatch, tmp_path):
    out = _dispatch(monkeypatch, tmp_path, n=3)  # <= threshold
    assert not list(out.glob("*.xlsx"))           # no attachment
    assert list(out.glob("consolidated__*.eml"))  # inline consolidated email


def test_sla_digest_attaches_full_list_when_many(monkeypatch, tmp_path):
    monkeypatch.setenv("COMMS_DRYRUN", "true")
    monkeypatch.setenv("COMMS_SENT_DIR", str(tmp_path / "sent"))
    monkeypatch.setenv("COMMS_AP_TEAM_MAILBOX", "ops@example.com")
    from app.config import get_settings
    get_settings.cache_clear()
    settings = get_settings()

    from app.comms import operator_notify

    breached = [
        {
            "invoice_id": f"INV-{i}", "vendor_name": "Acme", "invoice_amount": 1000.0 + i,
            "resolution_path": "ESCALATE_CONTROLLER", "sla_hours": 8, "status": "breached",
            "hours_overdue": 5.0, "comms_sent": False, "priority": "HIGH",
        }
        for i in range(1, 13)  # 12 > threshold of 10
    ]
    report = {
        "total_open": 12, "high_waiting": 12, "breached": 12, "due_soon": 0,
        "breached_items": breached, "due_soon_items": [],
    }
    res = operator_notify.send_sla_digest(run_id="run_dig", report=report, settings=settings)
    assert res["status"] == "dryrun"

    out = tmp_path / "sent" / "run_dig"
    xlsx = list(out.glob("sla-digest-*.xlsx"))
    assert len(xlsx) == 1  # full SLA list attached

    # The "recommended approach" column carries the friendly label, not the raw
    # routing code — consistent with the digest body and the reviewer console.
    from openpyxl import load_workbook

    ws = load_workbook(xlsx[0]).active
    header = [c.value for c in ws[1]]
    assert "Recommended approach" in header
    assert "Resolution path" not in header
    col = header.index("Recommended approach") + 1
    assert ws.cell(row=2, column=col).value == "Escalate to Controller"

    # With the full list attached, the email body is just the overview + a
    # pointer to the attachment — it does NOT re-list the invoices inline.
    eml = (out / next(p.name for p in out.glob("*.eml"))).read_text(encoding="utf-8", errors="ignore")
    assert "Full SLA list (12 items) attached" in eml
    assert "Breached (most overdue first):" not in eml
    assert "INV-1 |" not in eml  # no inline invoice rows
    get_settings.cache_clear()


def test_sla_digest_body_inlines_breakdown_only_without_attachment():
    """Under the attachment threshold the digest body inlines the most-overdue
    breakdown; once the spreadsheet is attached the body drops to the overview."""
    from app.comms.operator_notify import _body

    report = {
        "total_open": 2, "high_waiting": 2, "breached": 2, "due_soon": 0,
        "breached_items": [
            {"invoice_id": "INV-1", "vendor_name": "Acme", "invoice_amount": 1000.0,
             "resolution_path": "ESCALATE_CONTROLLER", "sla_hours": 8,
             "hours_overdue": 5.0},
        ],
        "due_soon_items": [],
    }
    with_list = _body("run_x", report, include_breakdown=True)
    assert "Breached (most overdue first):" in with_list
    assert "INV-1 |" in with_list

    overview = _body("run_x", report, include_breakdown=False)
    assert "Breached (most overdue first):" not in overview
    assert "INV-1 |" not in overview
    assert "Past SLA (breached): 2" in overview  # overview counts still present


def test_resolution_path_label():
    from app.schemas import resolution_path_label
    from app.schemas.resolution import ResolutionPath

    assert resolution_path_label("ESCALATE_CONTROLLER") == "Escalate to Controller"
    assert resolution_path_label(ResolutionPath.AUTO_APPROVE) == "Auto Approve"
    assert resolution_path_label("VENDOR_VALIDATION_REVIEW") == "Vendor Validation Review"
    assert resolution_path_label(None) == ""
    assert resolution_path_label("") == ""
    # Unknown/new code is humanized, never blank.
    assert resolution_path_label("SOME_NEW_PATH") == "Some New Path"
